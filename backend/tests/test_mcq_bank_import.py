"""MCQ Excel bank import workflow tests."""

from io import BytesIO
from uuid import UUID, uuid4

import pytest
from httpx import AsyncClient
from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import selectinload

from app.content.mcq_workbook import MCQ_COLUMNS, SCHEMA_VERSION, SHEET_QUESTIONS, SHEET_SETTINGS, SHEET_TAXONOMY
from app.content.mcq_workbook_template import build_template_workbook
from app.core.exceptions import AppException
from app.db.session import AsyncSessionLocal
from app.models.enums import McqImportMode
from app.models.mcq_bank import McqImportRun
from app.models.question import Question
from app.models.user import User
from app.services.mcq_bank_service import McqBankService


def _row(**kwargs):
    base = {
        "content_key": f"test.mcq.{uuid4().hex[:8]}",
        "bank_key": "test-bank",
        "domain_slug": "placement",
        "category_slug": "aptitude",
        "topic_slug": "percentages",
        "difficulty": "easy",
        "question_type": "single_choice",
        "question_text": "Test question?",
        "option_a": "A1",
        "option_b": "B1",
        "option_c": "C1",
        "option_d": "D1",
        "correct_option": "A",
        "explanation_a": "Because A.",
        "explanation_b": "Not B.",
        "explanation_c": "Not C.",
        "explanation_d": "Not D.",
        "solution_explanation": "Pick A.",
        "hint": "",
        "marks": 1,
        "negative_marks": 0,
        "estimated_time_seconds": 60,
        "status": "published",
    }
    base.update(kwargs)
    return base


def _workbook_bytes(rows: list[dict], bank_key: str = "test-bank") -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    settings = wb.create_sheet(SHEET_SETTINGS)
    settings.append(["schema_version", SCHEMA_VERSION])
    settings.append(["intended_bank_key", bank_key])
    tax = wb.create_sheet(SHEET_TAXONOMY)
    tax.append(["domain_slug", "category_slug", "topic_slug"])
    tax.append(["placement", "aptitude", "percentages"])
    q = wb.create_sheet(SHEET_QUESTIONS)
    q.append(MCQ_COLUMNS)
    for row in rows:
        q.append([row.get(col, "") for col in MCQ_COLUMNS])
    wb.create_sheet("Instructions")
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_student_cannot_stage_import(client: AsyncClient, student_auth):
    data = build_template_workbook("test-bank")
    files = {"file": ("t.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    response = await client.post(
        "/api/v1/admin/mcq-bank/import/stage",
        data={"bank_key": "test-bank", "import_mode": "add_update"},
        files=files,
        headers=student_auth[0],
    )
    assert response.status_code in (401, 403)


@pytest.mark.asyncio
async def test_import_preview_publish_and_idempotent(admin_auth):
    key = f"test.mcq.{uuid4().hex[:8]}"
    bank = f"test-bank-{uuid4().hex[:6]}"
    row = _row(content_key=key, bank_key=bank)
    payload = _workbook_bytes([row], bank_key=bank)

    async with AsyncSessionLocal() as db:
        service = McqBankService(db)
        from app.models.user import User

        admin = (await db.execute(select(User).where(User.email == "admin@jobready.dev"))).scalar_one()
        staged = await service.stage_upload(
            admin, bank_key=bank, import_mode=McqImportMode.ADD_UPDATE, data=payload
        )
        assert staged["can_publish"] is True
        run_id = staged["import_run_id"]
        published = await service.publish(admin, UUID(run_id))
        assert published["bank_version"] == 1

        question = (
            await db.execute(
                select(Question).options(selectinload(Question.options)).where(Question.content_key == key)
            )
        ).scalar_one()
        assert question.options[0].explanation == "Because A."

        staged2 = await service.stage_upload(
            admin, bank_key=bank, import_mode=McqImportMode.ADD_UPDATE, data=payload
        )
        assert len(staged2["preview"]["added"]) == 0
        assert len(staged2["preview"]["updated"]) == 0
        assert len(staged2["preview"]["unchanged"]) == 1


@pytest.mark.asyncio
async def test_invalid_row_blocks_publish(admin_auth):
    bank = f"test-bank-{uuid4().hex[:6]}"
    row = _row(bank_key=bank, correct_option="Z")
    payload = _workbook_bytes([row], bank_key=bank)
    async with AsyncSessionLocal() as db:
        service = McqBankService(db)
        admin = (await db.execute(select(User).where(User.email == "admin@jobready.dev"))).scalar_one()
        staged = await service.stage_upload(admin, bank_key=bank, import_mode=McqImportMode.ADD_UPDATE, data=payload)
        assert staged["can_publish"] is False
        run = (await db.execute(select(McqImportRun).where(McqImportRun.id == staged["import_run_id"]))).scalar_one()
        with pytest.raises(AppException):
            await service.publish(admin, run.id)
