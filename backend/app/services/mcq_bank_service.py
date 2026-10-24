"""Admin MCQ Excel bank: stage, preview, publish, export, rollback."""

from __future__ import annotations

import uuid
from datetime import UTC, datetime
from io import BytesIO
from typing import Any
from uuid import UUID, uuid4

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.content.mcq_workbook import (
    MCQ_COLUMNS,
    SCHEMA_VERSION,
    SHEET_INSTRUCTIONS,
    SHEET_QUESTIONS,
    SHEET_SETTINGS,
    SHEET_TAXONOMY,
    WorkbookParseResult,
    parse_workbook,
    row_content_hash,
)
from app.core.exceptions import AppException
from app.models.enums import Difficulty, McqImportMode, McqPublicationStatus, QuestionType
from app.models.mcq_bank import McqBankRegistry, McqImportRun, McqQuestionRevision
from app.models.question import Question, QuestionOption
from app.models.user import User
from app.repositories.question_repository import TaxonomyRepository
from app.services.catalog_service import CatalogService
from app.services.practice_service import question_snapshot
from app.utils.redis import get_redis


def _snapshot_for_revision(question: Question) -> dict:
    snap = question_snapshot(question)
    snap["content_key"] = question.content_key
    snap["bank_key"] = question.bank_key
    snap["hint"] = question.hint
    snap["publication_status"] = question.publication_status
    snap["revision"] = question.revision
    snap["domain_id"] = str(question.domain_id)
    snap["category_id"] = str(question.category_id)
    snap["topic_id"] = str(question.topic_id)
    for opt, item in zip(question.options, snap["options"], strict=False):
        item["explanation"] = opt.explanation
    return snap


class McqBankService:
    BANK_VERSION_CACHE_KEY = "practice:mcq_bank_version"

    def __init__(self, db: AsyncSession):
        self.db = db
        self.taxonomy = TaxonomyRepository(db)

    async def list_taxonomy_rows(self) -> list[dict[str, str]]:
        index = await self._taxonomy_index()
        return [
            {"domain_slug": d, "category_slug": c, "topic_slug": t}
            for (d, c, t) in sorted(index.keys())
        ]

    async def _taxonomy_index(self) -> dict[tuple[str, str, str], bool]:
        domains = await self.taxonomy.get_full_catalog()
        index: dict[tuple[str, str, str], bool] = {}
        for domain in domains:
            if not domain.is_active:
                continue
            for category in domain.categories:
                if not category.is_active:
                    continue
                for topic in category.topics:
                    if topic.is_active:
                        index[(domain.slug, category.slug, topic.slug)] = True
        return index

    async def _resolve_taxonomy(self, domain_slug: str, category_slug: str, topic_slug: str):
        domains = await self.taxonomy.get_full_catalog()
        for domain in domains:
            if domain.slug != domain_slug:
                continue
            for category in domain.categories:
                if category.slug != category_slug:
                    continue
                for topic in category.topics:
                    if topic.slug == topic_slug:
                        return domain, category, topic
        raise AppException(
            f"Unknown taxonomy: {domain_slug}/{category_slug}/{topic_slug}",
            status_code=400,
        )

    async def get_bank_meta(self, bank_key: str) -> dict:
        reg = await self._get_registry(bank_key)
        return {
            "bank_key": bank_key,
            "version": reg.version,
            "question_count": reg.question_count,
            "last_published_at": reg.last_published_at.isoformat() if reg.last_published_at else None,
        }

    async def list_banks(self) -> list[dict]:
        rows = (await self.db.execute(select(McqBankRegistry).order_by(McqBankRegistry.bank_key))).scalars().all()
        return [
            {
                "bank_key": row.bank_key,
                "version": row.version,
                "question_count": row.question_count,
                "last_published_at": row.last_published_at.isoformat() if row.last_published_at else None,
            }
            for row in rows
        ]

    async def stage_upload(
        self,
        admin: User,
        *,
        bank_key: str,
        import_mode: McqImportMode,
        data: bytes,
        allow_empty_replace: bool = False,
    ) -> dict:
        tax_index = await self._taxonomy_index()
        parsed = parse_workbook(data, taxonomy_index=tax_index)
        if parsed.settings.get("intended_bank_key") and parsed.settings["intended_bank_key"] != bank_key:
            raise AppException("Workbook intended_bank_key does not match selected bank", status_code=400)

        reg = await self._get_registry(bank_key)
        preview = await self._build_preview(
            bank_key=bank_key,
            import_mode=import_mode,
            parsed=parsed,
            allow_empty_replace=allow_empty_replace,
        )

        run = McqImportRun(
            id=uuid4(),
            bank_key=bank_key,
            import_mode=import_mode.value,
            status="staged",
            workbook_checksum=parsed.workbook_checksum,
            bank_version_at_preview=reg.version,
            staged_payload={"rows": [r.data for r in parsed.rows], "checksum": parsed.workbook_checksum},
            preview_summary=preview,
            row_errors=[e.as_dict() for e in parsed.errors],
            allow_empty_replace=allow_empty_replace,
            created_by=admin.id,
        )
        self.db.add(run)
        await self.db.commit()

        return {
            "import_run_id": str(run.id),
            "bank_key": bank_key,
            "import_mode": import_mode.value,
            "workbook_checksum": parsed.workbook_checksum,
            "bank_version_at_preview": reg.version,
            "can_publish": preview["can_publish"],
            "preview": preview,
            "errors": run.row_errors,
        }

    async def publish(self, admin: User, import_run_id: UUID) -> dict:
        run = await self._get_run(import_run_id)
        if run.status != "staged":
            raise AppException("Import run is not staged", status_code=400)
        if run.row_errors:
            raise AppException("Cannot publish: workbook has validation errors", status_code=400)

        reg = await self._get_registry(run.bank_key)
        if reg.version != run.bank_version_at_preview:
            raise AppException(
                "Bank changed since preview — re-upload and refresh preview",
                status_code=409,
            )

        preview = run.preview_summary
        if not preview.get("can_publish"):
            raise AppException("Preview marked this import as not publishable", status_code=400)

        mode = McqImportMode(run.import_mode)
        rows = run.staged_payload.get("rows", [])

        try:
            for item in preview.get("added", []):
                await self._upsert_row(run, item["content_key"], rows, creating=True)
            for item in preview.get("updated", []):
                await self._upsert_row(run, item["content_key"], rows, creating=False)
            if mode == McqImportMode.REPLACE_BANK:
                for item in preview.get("archived", []):
                    await self._archive_content_key(item["content_key"], run.bank_key)

            reg.version += 1
            reg.last_published_at = datetime.now(UTC)
            reg.question_count = await self._count_bank_questions(run.bank_key)
            run.status = "published"
            run.published_at = reg.last_published_at
            await self.db.commit()
        except Exception:
            await self.db.rollback()
            raise

        await CatalogService(self.db).invalidate_cache()
        await self._bump_bank_version_cache(reg.version)

        return {
            "import_run_id": str(run.id),
            "bank_key": run.bank_key,
            "bank_version": reg.version,
            "published_at": reg.last_published_at.isoformat(),
            "summary": preview,
        }

    async def rollback(self, admin: User, import_run_id: UUID) -> dict:
        run = await self._get_run(import_run_id)
        if run.status != "published":
            raise AppException("Only published imports can be rolled back", status_code=400)

        revisions = (
            await self.db.execute(
                select(McqQuestionRevision)
                .where(McqQuestionRevision.import_run_id == run.id)
                .order_by(McqQuestionRevision.created_at.desc())
            )
        ).scalars().all()

        restored = 0
        for rev in revisions:
            question = await self._get_question_by_id(rev.question_id)
            if question is None:
                continue
            await self._apply_snapshot(question, rev.snapshot, import_run_id=run.id)
            restored += 1

        reg = await self._get_registry(run.bank_key)
        reg.version += 1
        reg.question_count = await self._count_bank_questions(run.bank_key)
        run.status = "rolled_back"
        await self.db.commit()
        await CatalogService(self.db).invalidate_cache()
        await self._bump_bank_version_cache(reg.version)

        return {"import_run_id": str(run.id), "restored_questions": restored, "bank_version": reg.version}

    async def list_import_history(self, bank_key: str | None = None, limit: int = 30) -> list[dict]:
        stmt = select(McqImportRun).order_by(McqImportRun.created_at.desc()).limit(limit)
        if bank_key:
            stmt = stmt.where(McqImportRun.bank_key == bank_key)
        rows = (await self.db.execute(stmt)).scalars().all()
        return [
            {
                "id": str(row.id),
                "bank_key": row.bank_key,
                "import_mode": row.import_mode,
                "status": row.status,
                "created_at": row.created_at.isoformat(),
                "published_at": row.published_at.isoformat() if row.published_at else None,
                "error_count": len(row.row_errors or []),
                "preview_summary": {
                    "added": len((row.preview_summary or {}).get("added", [])),
                    "updated": len((row.preview_summary or {}).get("updated", [])),
                    "unchanged": len((row.preview_summary or {}).get("unchanged", [])),
                    "archived": len((row.preview_summary or {}).get("archived", [])),
                },
            }
            for row in rows
        ]

    async def export_bank(self, bank_key: str) -> bytes:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.worksheet.datavalidation import DataValidation

        questions = (
            await self.db.execute(
                select(Question)
                .options(selectinload(Question.options))
                .where(Question.bank_key == bank_key)
                .order_by(Question.content_key)
            )
        ).scalars().all()

        wb = Workbook()
        wb.remove(wb.active)
        settings = wb.create_sheet(SHEET_SETTINGS)
        settings.append(["schema_version", SCHEMA_VERSION])
        settings.append(["intended_bank_key", bank_key])

        tax_ws = wb.create_sheet(SHEET_TAXONOMY)
        tax_ws.append(["domain_slug", "category_slug", "topic_slug"])
        seen_tax: set[tuple[str, str, str]] = set()
        domains = await self.taxonomy.get_full_catalog()
        slug_map = {d.id: d for d in domains}
        for q in questions:
            dom = slug_map.get(q.domain_id)
            if not dom:
                continue
            cat = next((c for c in dom.categories if c.id == q.category_id), None)
            top = next((t for t in cat.topics if t.id == q.topic_id), None) if cat else None
            if not cat or not top:
                continue
            triple = (dom.slug, cat.slug, top.slug)
            if triple not in seen_tax:
                seen_tax.add(triple)
                tax_ws.append(list(triple))

        q_ws = wb.create_sheet(SHEET_QUESTIONS)
        q_ws.append(MCQ_COLUMNS)
        letter_map = ["A", "B", "C", "D"]
        for q in questions:
            opts = sorted(q.options, key=lambda o: o.sort_order)[:4]
            while len(opts) < 4:
                opts.append(QuestionOption(option_text="", is_correct=False, sort_order=len(opts)))
            correct_letter = "A"
            for i, opt in enumerate(opts):
                if opt.is_correct:
                    correct_letter = letter_map[i]
            status = q.publication_status or (
                McqPublicationStatus.PUBLISHED.value if q.is_active else McqPublicationStatus.ARCHIVED.value
            )
            dom = slug_map.get(q.domain_id)
            cat = next((c for c in dom.categories if c.id == q.category_id), None) if dom else None
            top = next((t for t in cat.topics if t.id == q.topic_id), None) if cat else None
            q_ws.append(
                [
                    q.content_key,
                    q.bank_key or bank_key,
                    dom.slug if dom else "",
                    cat.slug if cat else "",
                    top.slug if top else "",
                    q.difficulty.value,
                    q.question_type.value,
                    q.question_text,
                    opts[0].option_text,
                    opts[1].option_text,
                    opts[2].option_text,
                    opts[3].option_text,
                    correct_letter,
                    opts[0].explanation or "",
                    opts[1].explanation or "",
                    opts[2].explanation or "",
                    opts[3].explanation or "",
                    q.explanation or "",
                    q.hint or "",
                    q.marks,
                    q.negative_marks,
                    q.estimated_time_seconds,
                    status,
                ]
            )

        instr = wb.create_sheet(SHEET_INSTRUCTIONS)
        instr.append(["See docs/MCQ_BANK_ADMIN.md and mcq_workbook_schema.json"])
        instr[1][0].font = Font(bold=True)

        dv_diff = DataValidation(type="list", formula1='"easy,medium,hard"', allow_blank=False)
        dv_status = DataValidation(type="list", formula1='"draft,published,archived"', allow_blank=False)
        dv_correct = DataValidation(type="list", formula1='"A,B,C,D"', allow_blank=False)
        q_ws.add_data_validation(dv_diff)
        q_ws.add_data_validation(dv_status)
        q_ws.add_data_validation(dv_correct)

        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def build_template_bytes(self, bank_key: str, taxonomy_rows: list[dict[str, str]]) -> bytes:
        from app.content.mcq_workbook_template import build_template_workbook

        return build_template_workbook(bank_key, taxonomy_rows)

    async def get_content_version(self) -> int:
        try:
            redis = await get_redis()
            raw = await redis.get(self.BANK_VERSION_CACHE_KEY)
            if raw:
                return int(raw)
        except Exception:
            pass
        return 0

    async def _bump_bank_version_cache(self, version: int) -> None:
        try:
            redis = await get_redis()
            await redis.set(self.BANK_VERSION_CACHE_KEY, str(version))
        except Exception:
            pass

    async def _get_registry(self, bank_key: str) -> McqBankRegistry:
        reg = (
            await self.db.execute(select(McqBankRegistry).where(McqBankRegistry.bank_key == bank_key))
        ).scalar_one_or_none()
        if reg is None:
            reg = McqBankRegistry(bank_key=bank_key, version=0, question_count=0)
            self.db.add(reg)
            await self.db.flush()
        return reg

    async def _get_run(self, run_id: UUID) -> McqImportRun:
        run = (await self.db.execute(select(McqImportRun).where(McqImportRun.id == run_id))).scalar_one_or_none()
        if run is None:
            raise AppException("Import run not found", status_code=404)
        return run

    async def _count_bank_questions(self, bank_key: str) -> int:
        from sqlalchemy import func

        return (
            await self.db.execute(
                select(func.count())
                .select_from(Question)
                .where(
                    Question.bank_key == bank_key,
                    Question.publication_status != McqPublicationStatus.ARCHIVED.value,
                    Question.is_active.is_(True),
                )
            )
        ).scalar_one()

    async def _build_preview(
        self,
        *,
        bank_key: str,
        import_mode: McqImportMode,
        parsed: WorkbookParseResult,
        allow_empty_replace: bool,
    ) -> dict:
        existing = (
            await self.db.execute(
                select(Question).options(selectinload(Question.options)).where(Question.bank_key == bank_key)
            )
        ).scalars().all()
        by_key = {q.content_key: q for q in existing if q.content_key}
        incoming_keys = {r.data["content_key"] for r in parsed.rows}

        added: list[dict] = []
        updated: list[dict] = []
        unchanged: list[dict] = []
        invalid_cross_bank: list[dict] = []

        for row in parsed.rows:
            key = row.data["content_key"]
            other = (
                await self.db.execute(
                    select(Question).where(Question.content_key == key, Question.bank_key != bank_key)
                )
            ).scalar_one_or_none()
            if other is not None:
                invalid_cross_bank.append({"content_key": key, "row": row.row_number})
                continue

            current = by_key.get(key)
            if current is None:
                global_q = (
                    await self.db.execute(select(Question).where(Question.content_key == key))
                ).scalar_one_or_none()
                if global_q and global_q.bank_key and global_q.bank_key != bank_key:
                    invalid_cross_bank.append({"content_key": key, "row": row.row_number})
                    continue
                added.append({"content_key": key, "row": row.row_number})
            else:
                if await self._row_matches_question(row.data, current):
                    unchanged.append({"content_key": key, "row": row.row_number})
                else:
                    updated.append({"content_key": key, "row": row.row_number})

        archived: list[dict] = []
        if import_mode == McqImportMode.REPLACE_BANK:
            for key, question in by_key.items():
                if key not in incoming_keys and question.publication_status != McqPublicationStatus.ARCHIVED.value:
                    archived.append({"content_key": key, "question_id": str(question.id)})

        can_publish = not parsed.errors and not invalid_cross_bank
        if import_mode == McqImportMode.REPLACE_BANK and not incoming_keys and not allow_empty_replace:
            can_publish = False

        return {
            "added": added,
            "updated": updated,
            "unchanged": unchanged,
            "archived": archived,
            "invalid_cross_bank": invalid_cross_bank,
            "can_publish": can_publish,
            "blocking_reasons": self._blocking_reasons(
                parsed, invalid_cross_bank, import_mode, incoming_keys, allow_empty_replace
            ),
        }

    def _blocking_reasons(
        self,
        parsed: WorkbookParseResult,
        cross_bank: list,
        mode: McqImportMode,
        incoming_keys: set[str],
        allow_empty_replace: bool,
    ) -> list[str]:
        reasons: list[str] = []
        if parsed.errors:
            reasons.append(f"{len(parsed.errors)} validation error(s)")
        if cross_bank:
            reasons.append(f"{len(cross_bank)} cross-bank content_key collision(s)")
        if mode == McqImportMode.REPLACE_BANK and not incoming_keys and not allow_empty_replace:
            reasons.append("Replace mode with empty workbook requires explicit allow_empty_replace")
        return reasons

    async def _row_matches_question(self, row: dict, question: Question) -> bool:
        return row_content_hash(row) == row_content_hash(await self._question_as_row(question))

    async def _question_as_row(self, question: Question) -> dict:
        domains = await self.taxonomy.get_full_catalog()
        dom = next((d for d in domains if d.id == question.domain_id), None)
        cat = next((c for c in dom.categories if c.id == question.category_id), None) if dom else None
        top = next((t for t in cat.topics if t.id == question.topic_id), None) if cat else None
        opts = sorted(question.options, key=lambda o: o.sort_order)[:4]
        letters = ["A", "B", "C", "D"]
        correct = "A"
        for i, opt in enumerate(opts):
            if opt.is_correct:
                correct = letters[i]
        return {
            "content_key": question.content_key,
            "bank_key": question.bank_key,
            "domain_slug": dom.slug if dom else "",
            "category_slug": cat.slug if cat else "",
            "topic_slug": top.slug if top else "",
            "difficulty": question.difficulty.value,
            "question_type": question.question_type.value,
            "question_text": question.question_text,
            "option_a": opts[0].option_text if len(opts) > 0 else "",
            "option_b": opts[1].option_text if len(opts) > 1 else "",
            "option_c": opts[2].option_text if len(opts) > 2 else "",
            "option_d": opts[3].option_text if len(opts) > 3 else "",
            "correct_option": correct,
            "explanation_a": opts[0].explanation if len(opts) > 0 else None,
            "explanation_b": opts[1].explanation if len(opts) > 1 else None,
            "explanation_c": opts[2].explanation if len(opts) > 2 else None,
            "explanation_d": opts[3].explanation if len(opts) > 3 else None,
            "solution_explanation": question.explanation,
            "hint": question.hint,
            "marks": question.marks,
            "negative_marks": question.negative_marks,
            "estimated_time_seconds": question.estimated_time_seconds,
            "status": question.publication_status
            or (McqPublicationStatus.PUBLISHED.value if question.is_active else McqPublicationStatus.ARCHIVED.value),
        }

    async def _upsert_row(self, run: McqImportRun, content_key: str, rows: list[dict], *, creating: bool) -> None:
        row = next(r for r in rows if r["content_key"] == content_key)
        domain, category, topic = await self._resolve_taxonomy(
            row["domain_slug"], row["category_slug"], row["topic_slug"]
        )
        question = (
            await self.db.execute(
                select(Question).options(selectinload(Question.options)).where(Question.content_key == content_key)
            )
        ).scalar_one_or_none()

        if question is None:
            question = Question(
                id=uuid4(),
                question_type=QuestionType.SINGLE_CHOICE,
                content_key=content_key,
                bank_key=run.bank_key,
                question_text=row["question_text"],
                explanation=row.get("solution_explanation"),
                hint=row.get("hint"),
                difficulty=Difficulty(row["difficulty"]),
                domain_id=domain.id,
                category_id=category.id,
                topic_id=topic.id,
                marks=float(row["marks"]),
                negative_marks=float(row["negative_marks"]),
                estimated_time_seconds=int(row["estimated_time_seconds"]),
                revision=1,
                is_sample=False,
            )
            self.db.add(question)
        else:
            if question.bank_key and question.bank_key != run.bank_key:
                raise AppException(f"content_key {content_key} belongs to another bank", status_code=409)
            await self._save_revision(question, run.id)
            question.revision = (question.revision or 1) + 1
            question.domain_id = domain.id
            question.category_id = category.id
            question.topic_id = topic.id

        self._apply_row_to_question(question, row, run.bank_key)
        await self.db.flush()
        await self._replace_options(question, row)

    def _apply_row_to_question(self, question: Question, row: dict, bank_key: str) -> None:
        status = row["status"]
        question.bank_key = bank_key
        question.question_text = row["question_text"]
        question.explanation = row.get("solution_explanation")
        question.hint = row.get("hint")
        question.difficulty = Difficulty(row["difficulty"])
        question.marks = float(row["marks"])
        question.negative_marks = float(row["negative_marks"])
        question.estimated_time_seconds = int(row["estimated_time_seconds"])
        question.publication_status = status
        question.is_active = status == McqPublicationStatus.PUBLISHED.value

    async def _replace_options(self, question: Question, row: dict) -> None:
        from sqlalchemy import delete

        letters = ["A", "B", "C", "D"]
        texts = [row[f"option_{l.lower()}"] for l in letters]
        expls = [row.get(f"explanation_{l.lower()}") for l in letters]
        correct = row["correct_option"].upper()
        await self.db.execute(
            delete(QuestionOption).where(QuestionOption.question_id == question.id)
        )
        for i, letter in enumerate(letters):
            self.db.add(
                QuestionOption(
                    id=uuid4(),
                    question_id=question.id,
                    option_text=texts[i],
                    explanation=expls[i],
                    is_correct=letter == correct,
                    sort_order=i,
                )
            )

    async def _archive_content_key(self, content_key: str, bank_key: str) -> None:
        question = (
            await self.db.execute(select(Question).where(Question.content_key == content_key, Question.bank_key == bank_key))
        ).scalar_one_or_none()
        if question is None:
            return
        await self._save_revision(question, None)
        question.publication_status = McqPublicationStatus.ARCHIVED.value
        question.is_active = False
        question.revision = (question.revision or 1) + 1

    async def _save_revision(self, question: Question, import_run_id: UUID | None) -> None:
        snap = _snapshot_for_revision(question)
        self.db.add(
            McqQuestionRevision(
                id=uuid4(),
                question_id=question.id,
                content_key=question.content_key or "",
                bank_key=question.bank_key,
                revision=question.revision or 1,
                snapshot=snap,
                import_run_id=import_run_id,
            )
        )

    async def _apply_snapshot(self, question: Question, snapshot: dict, import_run_id: UUID) -> None:
        question.title = snapshot.get("title")
        question.question_text = snapshot["question_text"]
        question.explanation = snapshot.get("explanation")
        question.hint = snapshot.get("hint")
        question.marks = snapshot.get("marks", question.marks)
        question.negative_marks = snapshot.get("negative_marks", question.negative_marks)
        question.publication_status = snapshot.get("publication_status")
        question.bank_key = snapshot.get("bank_key")
        question.revision = (question.revision or 1) + 1
        question.is_active = question.publication_status == McqPublicationStatus.PUBLISHED.value
        for opt in list(question.options):
            await self.db.delete(opt)
        question.options.clear()
        for i, item in enumerate(snapshot.get("options", [])):
            question.options.append(
                QuestionOption(
                    id=UUID(item["id"]) if item.get("id") else uuid4(),
                    option_text=item["option_text"],
                    explanation=item.get("explanation"),
                    is_correct=item.get("is_correct", False),
                    sort_order=item.get("sort_order", i),
                )
            )

    async def _get_question_by_id(self, question_id: UUID) -> Question | None:
        return (
            await self.db.execute(
                select(Question).options(selectinload(Question.options)).where(Question.id == question_id)
            )
        ).scalar_one_or_none()
