"""Parse and validate JobReady MCQ Excel workbooks (schema v1)."""

from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass, field
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

SCHEMA_VERSION = "1.0"
SHEET_SETTINGS = "Settings"
SHEET_TAXONOMY = "Taxonomy"
SHEET_QUESTIONS = "MCQ_Questions"
SHEET_INSTRUCTIONS = "Instructions"

MCQ_COLUMNS = [
    "content_key",
    "bank_key",
    "domain_slug",
    "category_slug",
    "topic_slug",
    "difficulty",
    "question_type",
    "question_text",
    "option_a",
    "option_b",
    "option_c",
    "option_d",
    "correct_option",
    "explanation_a",
    "explanation_b",
    "explanation_c",
    "explanation_d",
    "solution_explanation",
    "hint",
    "marks",
    "negative_marks",
    "estimated_time_seconds",
    "status",
]

DIFFICULTIES = frozenset({"easy", "medium", "hard"})
STATUSES = frozenset({"draft", "published", "archived"})
CORRECT_LETTERS = frozenset({"A", "B", "C", "D"})
CONTENT_KEY_RE = re.compile(r"^[a-z0-9][a-z0-9._-]{2,159}$")
BANK_KEY_RE = re.compile(r"^[a-z0-9][a-z0-9._-]{1,79}$")

MARKS_MIN, MARKS_MAX = 0.0, 100.0
NEG_MARKS_MIN, NEG_MARKS_MAX = 0.0, 100.0
TIME_MIN, TIME_MAX = 5, 3600


@dataclass
class RowError:
    sheet: str
    row: int
    column: str
    message: str

    def as_dict(self) -> dict[str, Any]:
        return {"sheet": self.sheet, "row": self.row, "column": self.column, "message": self.message}


@dataclass
class ParsedMcqRow:
    row_number: int
    data: dict[str, Any]
    content_hash: str


@dataclass
class WorkbookParseResult:
    settings: dict[str, str]
    taxonomy_rows: list[dict[str, str]]
    rows: list[ParsedMcqRow] = field(default_factory=list)
    errors: list[RowError] = field(default_factory=list)
    workbook_checksum: str = ""

    @property
    def has_blocking_errors(self) -> bool:
        return bool(self.errors)


def workbook_checksum(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def row_content_hash(data: dict[str, Any]) -> str:
    canonical = json.dumps(data, sort_keys=True, ensure_ascii=False)
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()


def _cell_str(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value)
    text = str(value).strip()
    return text if text else None


def _read_settings(ws: Worksheet) -> dict[str, str]:
    out: dict[str, str] = {}
    for row in ws.iter_rows(min_row=1, max_col=2, values_only=True):
        key, val = row[0], row[1] if len(row) > 1 else None
        if key is None:
            continue
        k = str(key).strip().lower()
        v = _cell_str(val)
        if k and v:
            out[k] = v
    return out


def _read_sheet_table(ws: Worksheet) -> tuple[list[str], list[tuple[int, list[Any]]]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
    body: list[tuple[int, list[Any]]] = []
    for idx, row in enumerate(rows[1:], start=2):
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        body.append((idx, list(row)))
    return headers, body


def _formula_errors(ws: Worksheet, sheet_name: str, headers: list[str]) -> list[RowError]:
    errors: list[RowError] = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        if all(c.value is None or str(c.value).strip() == "" for c in row):
            continue
        for col_idx, cell in enumerate(row):
            if cell.data_type == "f":
                col = headers[col_idx] if col_idx < len(headers) else f"col_{col_idx + 1}"
                errors.append(
                    RowError(sheet_name, row_idx, col, "Formulas are not allowed in content cells")
                )
    return errors


def parse_workbook(
    data: bytes,
    *,
    taxonomy_index: dict[tuple[str, str, str], bool] | None = None,
) -> WorkbookParseResult:
    """Parse workbook bytes. taxonomy_index keys are (domain, category, topic) slugs."""
    checksum = workbook_checksum(data)
    result = WorkbookParseResult(settings={}, taxonomy_rows=[], workbook_checksum=checksum)

    try:
        wb = load_workbook(BytesIO(data), read_only=False, data_only=True, keep_links=False)
    except Exception as exc:
        result.errors.append(RowError("", 0, "", f"Unsupported or corrupt workbook: {exc}"))
        return result

    if wb.vba_archive is not None:
        result.errors.append(RowError("", 0, "", "Macro-enabled workbooks are not supported"))

    required = {SHEET_SETTINGS, SHEET_TAXONOMY, SHEET_QUESTIONS, SHEET_INSTRUCTIONS}
    missing = required - set(wb.sheetnames)
    for name in sorted(missing):
        result.errors.append(RowError(name, 0, "", "Required sheet is missing"))

    if missing:
        return result

    result.settings = _read_settings(wb[SHEET_SETTINGS])
    schema = result.settings.get("schema_version")
    if schema != SCHEMA_VERSION:
        result.errors.append(
            RowError(
                SHEET_SETTINGS,
                0,
                "schema_version",
                f"Expected schema_version {SCHEMA_VERSION}, got {schema!r}",
            )
        )

    tax_ws = wb[SHEET_TAXONOMY]
    tax_headers, tax_body = _read_sheet_table(tax_ws)
    for row_num, cells in tax_body:
        row_map = {
            tax_headers[i]: _cell_str(cells[i]) if i < len(cells) else None
            for i in range(len(tax_headers))
            if tax_headers[i]
        }
        if any(row_map.values()):
            result.taxonomy_rows.append(
                {
                    "domain_slug": row_map.get("domain_slug") or "",
                    "category_slug": row_map.get("category_slug") or "",
                    "topic_slug": row_map.get("topic_slug") or "",
                }
            )

    workbook_taxonomy = {
        (r["domain_slug"], r["category_slug"], r["topic_slug"])
        for r in result.taxonomy_rows
        if r["domain_slug"] and r["category_slug"] and r["topic_slug"]
    }

    q_ws = wb[SHEET_QUESTIONS]
    headers, body = _read_sheet_table(q_ws)
    header_set = set(headers)
    for col in MCQ_COLUMNS:
        if col not in header_set:
            result.errors.append(RowError(SHEET_QUESTIONS, 1, col, "Required column is missing"))

    result.errors.extend(_formula_errors(q_ws, SHEET_QUESTIONS, headers))

    seen_keys: set[str] = set()
    intended_bank = result.settings.get("intended_bank_key")

    for row_num, cells in body:
        row_data: dict[str, Any] = {}
        for i, col in enumerate(headers):
            if not col or col not in MCQ_COLUMNS:
                continue
            raw = cells[i] if i < len(cells) else None
            if col in {"marks", "negative_marks"}:
                row_data[col] = float(raw) if raw is not None and str(raw).strip() != "" else None
            elif col == "estimated_time_seconds":
                row_data[col] = int(raw) if raw is not None and str(raw).strip() != "" else None
            else:
                row_data[col] = _cell_str(raw)

        errors_for_row = _validate_row(
            row_data,
            row_num,
            intended_bank=intended_bank,
            workbook_taxonomy=workbook_taxonomy,
            taxonomy_index=taxonomy_index,
            seen_keys=seen_keys,
        )
        result.errors.extend(errors_for_row)
        if not errors_for_row:
            content_key = row_data["content_key"]
            seen_keys.add(content_key)
            normalized = _normalize_row(row_data)
            result.rows.append(
                ParsedMcqRow(
                    row_number=row_num,
                    data=normalized,
                    content_hash=row_content_hash(normalized),
                )
            )

    return result


def _normalize_row(data: dict[str, Any]) -> dict[str, Any]:
    out = dict(data)
    out["correct_option"] = str(out["correct_option"]).upper()
    out["question_type"] = str(out["question_type"]).lower()
    out["difficulty"] = str(out["difficulty"]).lower()
    out["status"] = str(out["status"]).lower()
    return out


def _validate_row(
    data: dict[str, Any],
    row_num: int,
    *,
    intended_bank: str | None,
    workbook_taxonomy: set[tuple[str, str, str]],
    taxonomy_index: dict[tuple[str, str, str], bool] | None,
    seen_keys: set[str],
) -> list[RowError]:
    errors: list[RowError] = []
    sheet = SHEET_QUESTIONS

    def err(col: str, msg: str) -> None:
        errors.append(RowError(sheet, row_num, col, msg))

    ck = data.get("content_key")
    if not ck:
        err("content_key", "content_key is required")
    elif not CONTENT_KEY_RE.match(ck):
        err("content_key", "content_key must be 3–160 chars: lowercase letters, digits, . _ -")

    bk = data.get("bank_key")
    if not bk:
        err("bank_key", "bank_key is required")
    elif not BANK_KEY_RE.match(bk):
        err("bank_key", "bank_key must be 2–80 chars: lowercase letters, digits, . _ -")
    elif intended_bank and bk != intended_bank:
        err("bank_key", f"bank_key must match Settings intended_bank_key ({intended_bank})")

    if ck and ck in seen_keys:
        err("content_key", "Duplicate content_key in workbook")

    for slug_col in ("domain_slug", "category_slug", "topic_slug"):
        if not data.get(slug_col):
            err(slug_col, f"{slug_col} is required")

    diff = (data.get("difficulty") or "").lower()
    if diff not in DIFFICULTIES:
        err("difficulty", "difficulty must be easy, medium, or hard")

    qtype = (data.get("question_type") or "").lower()
    if qtype != "single_choice":
        err("question_type", "Only single_choice is supported in schema v1")

    if not data.get("question_text"):
        err("question_text", "question_text is required")

    options = {
        "A": data.get("option_a"),
        "B": data.get("option_b"),
        "C": data.get("option_c"),
        "D": data.get("option_d"),
    }
    texts = [t for t in options.values() if t]
    if len(texts) != 4:
        err("option_a", "Exactly four non-empty options (A–D) are required")
    elif len(set(texts)) != 4:
        err("option_a", "All four options must be distinct")

    letter = (data.get("correct_option") or "").upper()
    if letter not in CORRECT_LETTERS:
        err("correct_option", "correct_option must be A, B, C, or D")
    elif letter and not options.get(letter):
        err("correct_option", f"correct_option {letter} has no option text")

    status = (data.get("status") or "").lower()
    if status not in STATUSES:
        err("status", "status must be draft, published, or archived")

    if status == "published":
        for letter_key, expl_col in zip(
            ("A", "B", "C", "D"),
            ("explanation_a", "explanation_b", "explanation_c", "explanation_d"),
            strict=True,
        ):
            if not data.get(expl_col):
                err(expl_col, f"{expl_col} is required when status is published")
        if not data.get("solution_explanation"):
            err("solution_explanation", "solution_explanation is required when status is published")

    marks = data.get("marks")
    if marks is None:
        err("marks", "marks is required")
    elif not MARKS_MIN <= float(marks) <= MARKS_MAX:
        err("marks", f"marks must be between {MARKS_MIN} and {MARKS_MAX}")

    neg = data.get("negative_marks")
    if neg is None:
        err("negative_marks", "negative_marks is required")
    elif not NEG_MARKS_MIN <= float(neg) <= NEG_MARKS_MAX:
        err("negative_marks", f"negative_marks must be between {NEG_MARKS_MIN} and {NEG_MARKS_MAX}")

    est = data.get("estimated_time_seconds")
    if est is None:
        err("estimated_time_seconds", "estimated_time_seconds is required")
    elif not TIME_MIN <= int(est) <= TIME_MAX:
        err("estimated_time_seconds", f"estimated_time_seconds must be between {TIME_MIN} and {TIME_MAX}")

    dom, cat, top = data.get("domain_slug"), data.get("category_slug"), data.get("topic_slug")
    if dom and cat and top:
        triple = (dom, cat, top)
        if triple not in workbook_taxonomy:
            err("topic_slug", "Taxonomy triple not listed on Taxonomy sheet")
        elif taxonomy_index is not None and not taxonomy_index.get(triple):
            err("topic_slug", "Taxonomy triple is not valid in the application database")

    return errors
