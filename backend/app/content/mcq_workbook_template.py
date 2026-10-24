"""Build downloadable MCQ Excel template (schema v1)."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from app.content.mcq_workbook import MCQ_COLUMNS, SCHEMA_VERSION, SHEET_INSTRUCTIONS, SHEET_QUESTIONS, SHEET_SETTINGS, SHEET_TAXONOMY

SAMPLE_ROWS = [
    {
        "content_key": "sample.aptitude.percent-discount",
        "bank_key": "phase11-aptitude",
        "domain_slug": "placement",
        "category_slug": "aptitude",
        "topic_slug": "percentages",
        "difficulty": "easy",
        "question_type": "single_choice",
        "question_text": "[SAMPLE] A ₹500 item has 20% off. What is the sale price?",
        "option_a": "₹400",
        "option_b": "₹100",
        "option_c": "₹450",
        "option_d": "₹480",
        "correct_option": "A",
        "explanation_a": "20% off ₹500 is ₹100 discount; ₹500 − ₹100 = ₹400.",
        "explanation_b": "₹100 is only the discount amount, not the final price.",
        "explanation_c": "₹450 would be a 10% discount, not 20%.",
        "explanation_d": "₹480 would be only a 4% discount.",
        "solution_explanation": "Discount = 0.20 × 500 = 100. Sale price = 500 − 100 = 400.",
        "hint": "Find the discount amount first, then subtract from the list price.",
        "marks": 1,
        "negative_marks": 0.25,
        "estimated_time_seconds": 60,
        "status": "draft",
    },
    {
        "content_key": "sample.technical.python-len",
        "bank_key": "phase11-technical",
        "domain_slug": "technical",
        "category_slug": "python",
        "topic_slug": "fundamentals",
        "difficulty": "easy",
        "question_type": "single_choice",
        "question_text": "[SAMPLE] What does len([1, 2, [3]]) return in Python 3?",
        "option_a": "3",
        "option_b": "4",
        "option_c": "2",
        "option_d": "TypeError",
        "correct_option": "A",
        "explanation_a": "len() counts top-level elements; the nested list counts as one element.",
        "explanation_b": "4 would count sub-elements inside the nested list, which len() does not do.",
        "explanation_c": "The list has three top-level items, not two.",
        "explanation_d": "Nested lists are valid; len() does not raise here.",
        "solution_explanation": "len([1, 2, [3]]) == 3 because the inner list is a single element.",
        "hint": "len() measures the number of items at the outermost level.",
        "marks": 1,
        "negative_marks": 0,
        "estimated_time_seconds": 45,
        "status": "draft",
    },
]


def build_template_workbook(bank_key: str, taxonomy_rows: list[dict[str, str]] | None = None) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    settings = wb.create_sheet(SHEET_SETTINGS)
    settings.append(["setting", "value"])
    settings.append(["schema_version", SCHEMA_VERSION])
    settings.append(["intended_bank_key", bank_key])

    tax = wb.create_sheet(SHEET_TAXONOMY)
    tax.append(["domain_slug", "category_slug", "topic_slug"])
    rows = taxonomy_rows or [
        {"domain_slug": "placement", "category_slug": "aptitude", "topic_slug": "percentages"},
        {"domain_slug": "technical", "category_slug": "python", "topic_slug": "fundamentals"},
    ]
    for row in rows:
        tax.append([row["domain_slug"], row["category_slug"], row["topic_slug"]])

    questions = wb.create_sheet(SHEET_QUESTIONS)
    questions.append(MCQ_COLUMNS)
    for sample in SAMPLE_ROWS:
        sample = dict(sample)
        sample["bank_key"] = bank_key
        questions.append([sample.get(col, "") for col in MCQ_COLUMNS])

    instr = wb.create_sheet(SHEET_INSTRUCTIONS)
    instr.append(["JobReady MCQ bank workbook — schema v1"])
    instr.append(["content_key", "Stable unique id; never reuse for a different question."])
    instr.append(["bank_key", "Must match Settings intended_bank_key and selected admin bank."])
    instr.append(["status", "draft | published | archived. Published rows require all explanations."])
    instr.append(["import", "Upload → validate → preview → publish. Local edits are not live until published."])
    for cell in instr["A"]:
        cell.font = Font(bold=cell.row == 1)

    dv_diff = DataValidation(type="list", formula1='"easy,medium,hard"')
    dv_status = DataValidation(type="list", formula1='"draft,published,archived"')
    dv_correct = DataValidation(type="list", formula1='"A,B,C,D"')
    dv_type = DataValidation(type="list", formula1='"single_choice"')
    for dv in (dv_diff, dv_status, dv_correct, dv_type):
        questions.add_data_validation(dv)
        dv.add(f"F2:F500")

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
